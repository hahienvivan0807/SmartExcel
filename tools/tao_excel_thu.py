"""Tạo folder excel/ – 5 dạng bảng MỚI, chưa từng dùng khi viết code, để kiểm thử "khách quan".

Bộ test_data/ dùng để phát triển: code đã được sửa cho tới khi qua hết → điểm cao là đương nhiên.
Folder excel/ giống "đề thi thử": cho biết chương trình xử lý bảng LẠ tốt tới đâu.

  1. ban_hang_quy.xlsx    Báo cáo kinh doanh: tiêu đề công ty, tiêu đề 2 tầng, ô gộp dọc,
                          dòng "Cộng" từng khu vực NẰM GIỮA bảng + dòng tổng cuối
  2. lop_hoc.xlsx         Cùng cấu trúc trên 3 sheet (10A1, 10A2, 10A3) – câu hỏi cần gộp nhiều sheet,
                          thông tin "lớp" chỉ nằm ở tên sheet
  3. thoi_khoa_bieu.xlsx  Bảng chéo mà ô bên trong là CHỮ (tên môn), có ô trống
  4. employees.xlsx       Bảng tiếng Anh, tên không dấu, lương "$1,200", ngày ISO, cột True/False
  5. nhap_xuat_kho.xlsx   HAI bảng trong CÙNG 1 sheet, cách nhau bởi dòng trống

Chạy: python tools/tao_excel_thu.py
"""
import json
import os
import shutil
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

GOC = os.path.join(os.path.dirname(__file__), "..", "excel")
XLSX = os.path.join(GOC, "xlsx_files")
shutil.rmtree(XLSX, ignore_errors=True)
os.makedirs(XLSX)

CAU_HOI, TANG1 = [], {}


def hoi(cau, dap_an, nhan, bang):
    CAU_HOI.append((cau, dap_an, nhan, bang))


def dep_trang(ws, dong_dam=()):
    """Cho file mở bằng Excel trông giống thật: chữ đậm tiêu đề, độ rộng cột vừa chữ."""
    for r in dong_dam:
        for c in ws[r]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
    from openpyxl.utils import get_column_letter
    for i, col in enumerate(ws.iter_cols(), 1):
        dai = max((len(str(c.value)) for c in col if c.value is not None and c.coordinate not in ws.merged_cells), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(dai + 2, 8), 40)


# =====================================================================
# 1. BÁO CÁO KINH DOANH – dòng "Cộng" từng nhóm nằm GIỮA bảng
# =====================================================================
khu_vuc = {
    "Miền Bắc": [("Nguyễn Văn Hùng", 120, 135, 150), ("Trần Thị Mai", 98, 110, 105)],
    "Miền Trung": [("Lê Văn Tâm", 80, 95, 88), ("Phạm Thị Hoa", 102, 99, 120), ("Hồ Minh Trí", 75, 82, 90)],
    "Miền Nam": [("Võ Thanh Sơn", 150, 160, 170), ("Đỗ Thị Ngọc", 130, 125, 140)],
}
wb = Workbook(); ws = wb.active; ws.title = "BaoCao"
ws.append(["CÔNG TY TNHH THƯƠNG MẠI XYZ"]); ws.merge_cells("A1:E1")
ws.append(["BÁO CÁO DOANH SỐ QUÝ 3/2026 (đơn vị: triệu đồng)"]); ws.merge_cells("A2:E2")
ws.append([])
ws.append(["Khu vực", "Nhân viên", "Doanh số", None, None])
ws.append([None, None, "Tháng 7", "Tháng 8", "Tháng 9"])
ws.merge_cells("A4:A5"); ws.merge_cells("B4:B5"); ws.merge_cells("C4:E4")
r = 6
for kv, ds in khu_vuc.items():
    for i, (ten, *so) in enumerate(ds):
        ws.append([kv if i == 0 else None, ten, *so])
    ws.merge_cells(f"A{r}:A{r + len(ds) - 1}")
    ws.append([None, f"Cộng {kv}", *[sum(d[k] for d in ds) for k in (1, 2, 3)]])
    r += len(ds) + 1
tat_ca = [d for ds in khu_vuc.values() for d in ds]
ws.append([None, "TỔNG CỘNG", *[sum(d[k] for d in tat_ca) for k in (1, 2, 3)]])
dep_trang(ws, dong_dam=(1, 2, 4, 5))
wb.save(os.path.join(XLSX, "ban_hang_quy.xlsx"))
B = "ban_hang_quy.xlsx::BaoCao"
TANG1[B] = {"huong": ["NGANG"], "so_dong": 7,
            "kieu": {"Khu vực": "chu", "Nhân viên": "chu", "Doanh số - Tháng 9": "so"},
            "o": [[1, "Khu vực", "Miền Bắc"], [2, "Nhân viên", "Lê Văn Tâm"]],
            "nhan": ["dong_cong_giua_bang", "tieu_de_2_tang", "o_gop_doc"]}
N = ["dong_cong_giua_bang"]
hoi("Tổng doanh số tháng 9 của toàn công ty là bao nhiêu?", sum(d[3] for d in tat_ca), N + ["tong"], [B])
hoi("Công ty có tất cả bao nhiêu nhân viên kinh doanh?", len(tat_ca), N + ["dem"], [B])
hoi("Nhân viên nào có tổng doanh số quý 3 cao nhất?", max(tat_ca, key=lambda d: sum(d[1:]))[0], N + ["max_tra_ten"], [B])
hoi("Doanh số tháng 8 của Phạm Thị Hoa là bao nhiêu?", 99, ["tra_o"], [B])

# =====================================================================
# 2. NHIỀU SHEET CÙNG CẤU TRÚC – "lớp" chỉ có ở tên sheet
# =====================================================================
lop = {
    "10A1": [("Nguyễn Hoàng Anh", "12/03/2010", 8.7, "Giỏi"), ("Bùi Thị Hằng", "05/07/2010", 7.2, "Khá"),
             ("Lê Quang Vinh", "21/11/2010", 6.4, "Trung bình"), ("Phan Thùy Linh", "02/01/2010", 9.1, "Giỏi")],
    "10A2": [("Trịnh Văn Long", "15/09/2010", 7.9, "Khá"), ("Mai Phương Thảo", "30/04/2010", 8.4, "Giỏi"),
             ("Cao Minh Khôi", "11/06/2010", 5.8, "Trung bình"), ("Lý Ngọc Hân", "19/12/2010", 8.9, "Giỏi"),
             ("Tạ Đức Huy", "08/08/2010", 7.0, "Khá")],
    "10A3": [("Đinh Công Mạnh", "25/02/2010", 6.9, "Trung bình"), ("Vương Thu Trang", "14/10/2010", 8.1, "Giỏi"),
             ("Kiều Anh Tuấn", "03/05/2010", 7.6, "Khá")],
}
wb = Workbook(); wb.remove(wb.active)
for ten_lop, hs in lop.items():
    ws = wb.create_sheet(ten_lop)
    ws.append(["STT", "Họ và tên", "Ngày sinh", "Điểm TB", "Xếp loại"])
    for i, h in enumerate(hs, 1):
        ws.append([i, *h])
    dep_trang(ws, dong_dam=(1,))
    TANG1[f"lop_hoc.xlsx::{ten_lop}"] = {
        "huong": ["NGANG"], "so_dong": len(hs),
        "kieu": {"Họ và tên": "chu", "Ngày sinh": "ngay", "Điểm TB": "so"}, "nhan": ["nhieu_sheet"]}
wb.save(os.path.join(XLSX, "lop_hoc.xlsx"))
BL = [f"lop_hoc.xlsx::{l}" for l in lop]
hoi("Toàn khối có bao nhiêu học sinh xếp loại Giỏi?",
    sum(h[3] == "Giỏi" for hs in lop.values() for h in hs), ["nhieu_sheet", "gop_sheet", "dem"], BL)
hoi("Điểm trung bình cao nhất của lớp 10A2 là bao nhiêu?", max(h[2] for h in lop["10A2"]),
    ["nhieu_sheet", "max"], ["lop_hoc.xlsx::10A2"])
hoi("Học sinh Đinh Công Mạnh học lớp nào?", "10A3", ["nhieu_sheet", "dap_an_la_ten_sheet"], ["lop_hoc.xlsx::10A3"])

# =====================================================================
# 3. BẢNG CHÉO Ô CHỮ – thời khóa biểu
# =====================================================================
thu = ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7"]
tkb = [  # mỗi dòng = 1 tiết
    ["Chào cờ", "Toán", "Văn", "Lý", "Toán", "Sinh"],
    ["Toán", "Văn", "Anh", "Hóa", "Văn", "Sinh"],
    ["Văn", "Anh", "Toán", "Sử", "Tin", "Thể dục"],
    ["Lý", "Hóa", "Địa", "Toán", "Anh", None],
    ["Anh", "Sinh hoạt", None, "GDCD", "Hóa", None],
]
wb = Workbook(); ws = wb.active; ws.title = "TKB"
ws.append(["THỜI KHÓA BIỂU LỚP 11B – HỌC KỲ I"]); ws.merge_cells("A1:G1")
ws.append(["Tiết", *thu])
for i, dong in enumerate(tkb, 1):
    ws.append([i, *dong])
dep_trang(ws, dong_dam=(1, 2))
wb.save(os.path.join(XLSX, "thoi_khoa_bieu.xlsx"))
B = "thoi_khoa_bieu.xlsx::TKB"
TANG1[B] = {"huong": ["NGANG", "CHEO"], "so_dong": 5,
            "cot": ["Tiết", *thu], "o": [[2, "Thứ 4", "Toán"]], "nhan": ["bang_cheo_o_chu"]}
C = ["bang_cheo_o_chu"]
hoi("Thứ 4 tiết 3 học môn gì?", tkb[2][2], C + ["tra_o"], [B])
hoi("Môn Toán có bao nhiêu tiết trong một tuần?", sum(o == "Toán" for d in tkb for o in d), C + ["dem_toan_bang"], [B])
hoi("Thứ 7 học mấy tiết?", sum(d[5] is not None for d in tkb), C + ["dem_o_trong"], [B])

# =====================================================================
# 4. BẢNG TIẾNG ANH – tên không dấu, "$1,200", ngày ISO dạng chữ, True/False
# =====================================================================
staff = [
    ("E001", "Ngo Thi Lan", "Sales", 1200, "2021-03-15", True),
    ("E002", "Tran Minh Duc", "IT", 1850, "2019-07-01", True),
    ("E003", "Le Van Hai", "HR", 1100, "2022-11-20", False),
    ("E004", "Pham Thu Trang", "Sales", 1350, "2020-01-06", True),
    ("E005", "Nguyen Quoc Bao", "IT", 2100, "2018-05-14", True),
    ("E006", "Ngo Van Nam", "Sales", 1500, "2023-09-04", False),
]
wb = Workbook(); ws = wb.active; ws.title = "Staff"
ws.append(["Employee ID", "Full Name", "Department", "Salary", "Start Date", "Active"])
for s in staff:
    ws.append([s[0], s[1], s[2], f"${s[3]:,}", s[4], s[5]])
dep_trang(ws, dong_dam=(1,))
wb.save(os.path.join(XLSX, "employees.xlsx"))
B = "employees.xlsx::Staff"
TANG1[B] = {"huong": ["NGANG"], "so_dong": 6,
            "kieu": {"Full Name": "chu", "Department": "chu", "Salary ($)": "so", "Start Date": "ngay"},
            "o": [[0, "Salary ($)", 1200]], "nhan": ["tieng_anh", "tien_do_la"]}
E = ["tieng_anh"]
hoi("Who has the highest salary?", max(staff, key=lambda s: s[3])[1], E + ["cau_hoi_tieng_anh", "tien_do_la", "max_tra_ten"], [B])
hoi("Phòng Sales có bao nhiêu nhân viên đang làm việc (Active)?",
    sum(s[2] == "Sales" and s[5] for s in staff), E + ["dung_sai", "nhieu_dieu_kien"], [B])
hoi("Tổng lương của phòng IT là bao nhiêu đô la?", sum(s[3] for s in staff if s[2] == "IT"),
    E + ["tien_do_la", "tong_dieu_kien"], [B])
hoi("Ai vào công ty sớm nhất?", min(staff, key=lambda s: s[4])[1], E + ["ngay_chu", "max_tra_ten"], [B])

# =====================================================================
# 5. HAI BẢNG TRONG 1 SHEET – nhập kho / xuất kho
# =====================================================================
nhap = [("01/10/2026", "Gạo", 500, "kg"), ("03/10/2026", "Đường", 200, "kg"), ("05/10/2026", "Gạo", 300, "kg"),
        ("08/10/2026", "Dầu ăn", 120, "chai"), ("12/10/2026", "Mì gói", 50, "thùng")]
xuat = [("04/10/2026", "Gạo", 250, "Siêu thị Minh Anh"), ("06/10/2026", "Đường", 80, "Tạp hóa Bảy Hiền"),
        ("09/10/2026", "Dầu ăn", 60, "Siêu thị Minh Anh"), ("13/10/2026", "Gạo", 150, "Quán cơm Hương Quê"),
        ("15/10/2026", "Mì gói", 20, "Tạp hóa Bảy Hiền")]
wb = Workbook(); ws = wb.active; ws.title = "NhapXuat"
ws.append(["SỔ THEO DÕI NHẬP – XUẤT KHO THÁNG 10/2026"]); ws.merge_cells("A1:D1")
ws.append([])
ws.append(["BẢNG NHẬP KHO"]); ws.merge_cells("A3:D3")
ws.append(["Ngày", "Mặt hàng", "Số lượng", "Đơn vị tính"])
for n in nhap:
    ws.append(list(n))
ws.append([]); ws.append([])
dong_xuat = ws.max_row + 1
ws.append(["BẢNG XUẤT KHO"]); ws.merge_cells(f"A{dong_xuat}:D{dong_xuat}")
ws.append(["Ngày", "Mặt hàng", "Số lượng", "Khách hàng"])
for x in xuat:
    ws.append(list(x))
dep_trang(ws, dong_dam=(1, 3, 4, dong_xuat, dong_xuat + 1))
wb.save(os.path.join(XLSX, "nhap_xuat_kho.xlsx"))
B1, B2 = "nhap_xuat_kho.xlsx::NhapXuat#1", "nhap_xuat_kho.xlsx::NhapXuat#2"
TANG1[B1] = {"huong": ["NGANG"], "so_dong": 5, "kieu": {"Ngày": "ngay", "Số lượng": "so"},
             "o": [[0, "Mặt hàng", "Gạo"]], "nhan": ["hai_bang_mot_sheet"]}
TANG1[B2] = {"huong": ["NGANG"], "so_dong": 5, "kieu": {"Ngày": "ngay", "Số lượng": "so", "Khách hàng": "chu"},
             "o": [[0, "Khách hàng", "Siêu thị Minh Anh"]], "nhan": ["hai_bang_mot_sheet"]}
H = ["hai_bang_mot_sheet"]
hoi("Tổng số lượng Gạo đã nhập kho là bao nhiêu?", sum(n[2] for n in nhap if n[1] == "Gạo"), H + ["tong_dieu_kien"], [B1])
hoi("Tổng số lượng Gạo đã xuất kho là bao nhiêu?", sum(x[2] for x in xuat if x[1] == "Gạo"), H + ["tong_dieu_kien"], [B2])
hoi("Siêu thị Minh Anh đã mua những mặt hàng nào?",
    ", ".join(x[1] for x in xuat if x[3] == "Siêu thị Minh Anh"), H + ["liet_ke"], [B2])


# =====================================================================
def dep(x):
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x)


queries, dap_an = [], {}
for i, (cau, da, nhan, bang) in enumerate(CAU_HOI, 1):
    queries.append({"id": i, "question": cau})
    dap_an[str(i)] = {"answer": dep(da), "nhan": nhan, "bang": bang}
for ten, du_lieu in [("queries.json", queries), ("dap_an.json", dap_an), ("tang1.json", TANG1)]:
    with open(os.path.join(GOC, ten), "w", encoding="utf-8") as f:
        json.dump(du_lieu, f, ensure_ascii=False, indent=2)
print(f"Đã tạo {len(os.listdir(XLSX))} file Excel ({len(TANG1)} bảng mong đợi), {len(queries)} câu hỏi trong {os.path.abspath(GOC)}")

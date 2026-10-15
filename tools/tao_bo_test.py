"""Tạo bộ test theo 3 trục: CẤU TRÚC bảng × KIỂU DỮ LIỆU × LOẠI CÂU HỎI.
Lĩnh vực (giáo dục, nhân sự, bán hàng, logistics, y tế, tài chính, kho) chỉ là "lớp áo" xoay vòng.

Tạo ra trong test_data/:
  xlsx_files/   các file Excel
  queries.json  câu hỏi (đúng dạng máy chấm: id + question)
  dap_an.json   đáp án đúng + nhãn + bảng cần dùng (cho tầng 2)
  tang1.json    kết quả làm sạch mong đợi của từng bảng (cho tầng 1)

Chạy: python tools/tao_bo_test.py
"""
import json
import os
import random
import shutil
from datetime import date, datetime, timedelta

from openpyxl import Workbook

GOC = os.path.join(os.path.dirname(__file__), "..", "test_data")
XLSX = os.path.join(GOC, "xlsx_files")
shutil.rmtree(XLSX, ignore_errors=True)
os.makedirs(XLSX)

CAU_HOI = []   # (câu hỏi, đáp án, [nhãn], [bảng cần dùng])
TANG1 = {}     # tên bảng → kỳ vọng sau Bước 1


def hoi(cau, dap_an, nhan, bang):
    CAU_HOI.append((cau, dap_an, nhan, bang))


def luu(wb, ten_file):
    wb.save(os.path.join(XLSX, ten_file))


def vn(x, nghin="."):
    """1200000 → '1.200.000'"""
    return f"{x:,}".replace(",", nghin)


# =====================================================================
# 1. GIÁO DỤC – ngang, tiêu đề báo cáo phía trên, dòng "Tổng cộng", số thập phân
# =====================================================================
hoc_sinh = [
    ("Nguyễn Văn An", "10A", "Nam", 8.5, 7.0),
    ("Trần Thị Bình", "10B", "Nữ", 9.0, 8.5),
    ("Ngô Minh Châu", "10A", "Nữ", 6.0, 9.0),
    ("Lê Hoàng Dũng", "10C", "Nam", 7.5, 6.5),
    ("Phạm Thu Hà", "10A", "Nữ", 9.5, 8.0),
    ("Võ Quốc Huy", "10B", "Nam", 5.0, 7.5),
    ("Đặng Mai Lan", "10C", "Nữ", 8.0, 9.5),
]
wb = Workbook(); ws = wb.active; ws.title = "HK1"
ws.append(["TRƯỜNG THPT ABC"]); ws.append(["BẢNG ĐIỂM HỌC KỲ I"]); ws.append([])
ws.append(["STT", "Họ tên", "Lớp", "Giới tính", "Toán", "Văn"])
for i, hs in enumerate(hoc_sinh, 1):
    ws.append([i, *hs])
ws.append(["", "Tổng cộng", "", "", sum(h[3] for h in hoc_sinh), sum(h[4] for h in hoc_sinh)])
luu(wb, "diem_thi.xlsx")
B = "diem_thi.xlsx::HK1"
TANG1[B] = {"huong": ["NGANG"], "so_dong": 7,
            "kieu": {"Họ tên": "chu", "Lớp": "chu", "Toán": "so", "Văn": "so"},
            "o": [[0, "Họ tên", "Nguyễn Văn An"], [6, "Văn", 9.5]],
            "nhan": ["tieu_de_bao_cao", "dong_tong", "so_thap_phan"]}
hoi("Lớp 10A có bao nhiêu học sinh?", 3, ["giao_duc", "dem_dieu_kien"], [B])
hoi("Ai có điểm Toán cao nhất?", "Phạm Thu Hà", ["giao_duc", "max_tra_ten"], [B])
hoi("Điểm Văn của Ngô Minh Châu là bao nhiêu?", 9, ["giao_duc", "tra_o"], [B])
hoi("Có bao nhiêu học sinh nữ?", 4, ["giao_duc", "dem_dieu_kien"], [B])
hoi("Điểm trung bình môn Toán của lớp 10A là bao nhiêu?", round((8.5 + 6.0 + 9.5) / 3, 2),
    ["giao_duc", "trung_binh", "lam_tron"], [B])

# =====================================================================
# 2. NHÂN SỰ – bảng DỌC (tên thuộc tính ở cột đầu)
# =====================================================================
nhan_vien = {
    "Họ tên": ["Hoàng Minh", "Lý Thu Trang", "Bùi Văn Tú", "Đỗ Ngọc Anh"],
    "Phòng ban": ["Kế toán", "Marketing", "Kế toán", "Kỹ thuật"],
    "Tuổi": [34, 28, 45, 31],
    "Lương": [15000000, 18000000, 22000000, 20000000],
}
wb = Workbook(); ws = wb.active; ws.title = "NhanVien"
for k, v in nhan_vien.items():
    ws.append([k, *v])
luu(wb, "nhan_vien.xlsx")
B = "nhan_vien.xlsx::NhanVien"
TANG1[B] = {"huong": ["DOC"], "so_dong": 4,
            "kieu": {"Họ tên": "chu", "Phòng ban": "chu", "Tuổi": "so", "Lương": "so"},
            "o": [[2, "Họ tên", "Bùi Văn Tú"], [2, "Lương", 22000000]],
            "nhan": ["bang_doc"]}
hoi("Phòng Kế toán có mấy nhân viên?", 2, ["nhan_su", "bang_doc", "dem_dieu_kien"], [B])
hoi("Nhân viên lớn tuổi nhất tên gì?", "Bùi Văn Tú", ["nhan_su", "bang_doc", "max_tra_ten"], [B])

# =====================================================================
# 3. BÁN HÀNG – bảng CHÉO (dòng = sản phẩm, cột = tháng)
# =====================================================================
thang = ["Tháng 1", "Tháng 2", "Tháng 3"]
doanh_thu = {"Áo thun": [120, 150, 90], "Quần jean": [200, 180, 210], "Giày": [80, 95, 130]}
wb = Workbook(); ws = wb.active; ws.title = "DoanhThu"
ws.append(["Sản phẩm", *thang])
for sp, so in doanh_thu.items():
    ws.append([sp, *so])
luu(wb, "doanh_thu.xlsx")
B = "doanh_thu.xlsx::DoanhThu"
TANG1[B] = {"huong": ["NGANG", "CHEO"], "so_dong": 3,
            "kieu": {"Sản phẩm": "chu", "Tháng 1": "so", "Tháng 3": "so"},
            "o": [[1, "Tháng 3", 210]], "nhan": ["bang_cheo"]}
hoi("Tổng doanh thu của Quần jean trong 3 tháng là bao nhiêu?", 590, ["ban_hang", "bang_cheo", "tong_dieu_kien"], [B])
hoi("Sản phẩm nào có doanh thu tháng 3 cao nhất?", "Quần jean", ["ban_hang", "bang_cheo", "max_tra_ten"], [B])

# =====================================================================
# 4. KINH DOANH – tiêu đề 2 TẦNG (ô gộp) + số dạng chữ "1.200.000"
# =====================================================================
chi_nhanh = {"Hà Nội": (1200000, 1500000), "Đà Nẵng": (800000, 950000), "TP.HCM": (2100000, 1900000)}
wb = Workbook(); ws = wb.active; ws.title = "DoanhSo"
ws.append(["Chi nhánh", "Doanh số", None]); ws.append([None, "Quý 1", "Quý 2"])
ws.merge_cells("A1:A2"); ws.merge_cells("B1:C1")
for ten, (q1, q2) in chi_nhanh.items():
    ws.append([ten, vn(q1), vn(q2)])
luu(wb, "chi_nhanh.xlsx")
B = "chi_nhanh.xlsx::DoanhSo"
TANG1[B] = {"huong": ["NGANG"], "so_dong": 3,
            "cot": ["Chi nhánh", "Doanh số - Quý 1", "Doanh số - Quý 2"],
            "kieu": {"Doanh số - Quý 1": "so", "Doanh số - Quý 2": "so"},
            "o": [[1, "Doanh số - Quý 1", 800000]],
            "nhan": ["tieu_de_2_tang", "so_phan_cach_nghin"]}
hoi("Doanh số quý 2 của chi nhánh Đà Nẵng là bao nhiêu?", 950000, ["kinh_doanh", "tieu_de_2_tang", "tra_o"], [B])
hoi("Chi nhánh nào có tổng doanh số 2 quý cao nhất?", "TP.HCM", ["kinh_doanh", "tieu_de_2_tang", "max_tra_ten"], [B])

# =====================================================================
# 5. LOGISTICS – tiêu đề gộp, ô gộp DỌC, SĐT/mã bưu chính có số 0 đầu,
#    số có đơn vị "12,5 kg", tiền "35.000 đ", ngày dạng chữ, giờ "08:30"
# =====================================================================
van_don = [  # mã, kho, tỉnh nhận, sđt, mã bưu chính, kg, ngày gửi, giờ, cước, trạng thái
    ("VD0001", "Bình Dương", "Hà Nội", "0901234567", "100000", 12.5, "01/10/2026", "08:30", 35000, "Đã giao"),
    ("VD0002", "Bình Dương", "Đà Nẵng", "0912345678", "550000", 3.2, "03/10/2026", "09:15", 28000, "Đã giao"),
    ("VD0003", "Bình Dương", "Cần Thơ", "0987654321", "900000", 25, "05/10/2026", "07:45", 52000, "Hoàn"),
    ("VD0004", "Bình Dương", "Hà Nội", "0933111222", "100000", 8, "28/09/2026", "10:00", 30000, "Đã giao"),
    ("VD0005", "Bình Dương", "Huế", "0944555666", "530000", 1.5, "12/10/2026", "14:20", 22000, "Đang giao"),
    ("VD0006", "Hà Nội", "Đà Nẵng", "0955777888", "550000", 40, "02/10/2026", "08:00", 75000, "Đã giao"),
    ("VD0007", "Hà Nội", "TP.HCM", "0966999000", "070000", 15.75, "07/10/2026", "11:30", 48000, "Hoàn"),
    ("VD0008", "Hà Nội", "Bình Dương", "0977123123", "075000", 6, "30/09/2026", "13:00", 26000, "Đã giao"),
    ("VD0009", "Hà Nội", "Hà Giang", "0988456456", "310000", 2.25, "14/10/2026", "15:45", 45000, "Đang giao"),
    ("VD0010", "Đà Nẵng", "Hà Nội", "0399888777", "100000", 10, "09/10/2026", "09:00", 40000, "Đã giao"),
    ("VD0011", "Đà Nẵng", "Quảng Nam", "0388111999", "560000", 30, "11/10/2026", "16:10", 36000, "Hoàn"),
    ("VD0012", "Đà Nẵng", "TP.HCM", "0377222333", "070000", 4.5, "25/09/2026", "10:30", 33000, "Đã giao"),
]
wb = Workbook(); ws = wb.active; ws.title = "VanDon"
ws.append(["DANH SÁCH VẬN ĐƠN THÁNG 10/2026"] + [None] * 10)
ws.merge_cells("A1:K1")
ws.append(["STT", "Mã vận đơn", "Kho xuất", "Tỉnh nhận", "SĐT người nhận", "Mã bưu chính",
           "Trọng lượng", "Ngày gửi", "Giờ xuất", "Cước phí", "Trạng thái"])
kho_truoc, dong_dau_kho = None, 3
for i, (ma, kho, tinh, sdt, mbc, kg, ngay, gio, cuoc, tt) in enumerate(van_don, 1):
    r = i + 2
    kg_chu = (f"{kg:g}".replace(".", ",")) + " kg"
    ws.append([i, ma, kho if kho != kho_truoc else None, tinh, sdt, mbc, kg_chu, ngay, gio, f"{vn(cuoc)} đ", tt])
    if kho != kho_truoc:
        if kho_truoc is not None and r - 1 > dong_dau_kho:
            ws.merge_cells(f"C{dong_dau_kho}:C{r - 1}")
        kho_truoc, dong_dau_kho = kho, r
ws.merge_cells(f"C{dong_dau_kho}:C{len(van_don) + 2}")
ws.append([None, "Tổng cộng", None, None, None, None,
           f"{sum(v[5] for v in van_don):g}".replace(".", ",") + " kg", None, None,
           f"{vn(sum(v[8] for v in van_don))} đ", None])
luu(wb, "van_don.xlsx")
B = "van_don.xlsx::VanDon"
TANG1[B] = {"huong": ["NGANG"], "so_dong": 12,
            "kieu": {"Mã vận đơn": "chu", "SĐT người nhận": "chu", "Mã bưu chính": "chu",
                     "Trọng lượng (kg)": "so", "Ngày gửi": "ngay", "Giờ xuất": "chu", "Cước phí (đ)": "so"},
            "o": [[0, "SĐT người nhận", "0901234567"], [6, "Mã bưu chính", "070000"],
                  [0, "Mã bưu chính", "100000"], [0, "Trọng lượng (kg)", 12.5],
                  [1, "Kho xuất", "Bình Dương"], [11, "Kho xuất", "Đà Nẵng"],
                  [0, "Ngày gửi", "2026-10-01"], [0, "Giờ xuất", "08:30"], [0, "Cước phí (đ)", 35000]],
            "nhan": ["tieu_de_gop", "o_gop_doc", "ma_so_0", "so_don_vi", "tien", "ngay_chu", "gio", "dong_tong"]}
L = ["logistics"]
hoi("Tổng trọng lượng các đơn gửi đi Hà Nội là bao nhiêu kg?",
    sum(v[5] for v in van_don if v[2] == "Hà Nội"), L + ["so_don_vi", "tong_dieu_kien"], [B])
hoi("Có bao nhiêu đơn được gửi trong tháng 10/2026?",
    sum(v[6].endswith("/10/2026") for v in van_don), L + ["ngay_chu", "dem_dieu_kien"], [B])
hoi("Số điện thoại 0399888777 là của mã vận đơn nào?", "VD0010", L + ["ma_so_0", "tra_o"], [B])
hoi("Kho Bình Dương có bao nhiêu đơn bị hoàn?",
    sum(v[1] == "Bình Dương" and v[9] == "Hoàn" for v in van_don), L + ["o_gop_doc", "nhieu_dieu_kien"], [B])
hoi("Liệt kê các mã vận đơn có mã bưu chính 070000.",
    ", ".join(v[0] for v in van_don if v[4] == "070000"), L + ["ma_so_0", "liet_ke"], [B])
hoi("Tổng cước phí của các đơn đã giao là bao nhiêu?",
    sum(v[8] for v in van_don if v[9] == "Đã giao"), L + ["tien", "tong_dieu_kien"], [B])
hoi("kho ha noi xuat bao nhieu don?", sum(v[1] == "Hà Nội" for v in van_don), L + ["khong_dau", "dem_dieu_kien"], [B])
hoi("Đơn nào nặng nhất?", max(van_don, key=lambda v: v[5])[0], L + ["so_don_vi", "max_tra_ten"], [B])

# =====================================================================
# 6. Y TẾ – 2 SHEET liên quan nhau (cần ghép), ngày Excel, phần trăm "80%"
# =====================================================================
benh_nhan = [
    ("BN001", "Trần Văn Bảo", date(1980, 5, 12), "Nam", "Nội"),
    ("BN002", "Lê Thị Cúc", date(1992, 11, 3), "Nữ", "Sản"),
    ("BN003", "Phạm Minh Đức", date(1975, 1, 20), "Nam", "Ngoại"),
    ("BN004", "Hoàng Thị Em", date(2001, 7, 30), "Nữ", "Nội"),
    ("BN005", "Vũ Văn Phát", date(1968, 3, 15), "Nam", "Tim mạch"),
]
kham = [
    ("BN001", date(2026, 9, 2), "Viêm họng", 250000, "80%"),
    ("BN003", date(2026, 9, 5), "Gãy tay", 1800000, "100%"),
    ("BN001", date(2026, 9, 20), "Tăng huyết áp", 400000, "80%"),
    ("BN002", date(2026, 10, 1), "Khám thai", 300000, "100%"),
    ("BN005", date(2026, 10, 3), "Rối loạn nhịp tim", 2500000, "80%"),
    ("BN001", date(2026, 10, 8), "Tái khám", 150000, "80%"),
    ("BN004", date(2026, 10, 10), "Cảm cúm", 200000, "0%"),
]
wb = Workbook(); ws = wb.active; ws.title = "BenhNhan"
ws.append(["Mã BN", "Họ tên", "Ngày sinh", "Giới tính", "Khoa"])
for bn in benh_nhan:
    ws.append([bn[0], bn[1], datetime.combine(bn[2], datetime.min.time()), bn[3], bn[4]])
ws2 = wb.create_sheet("KhamBenh")
ws2.append(["Mã BN", "Ngày khám", "Chẩn đoán", "Chi phí", "Tỷ lệ BHYT"])
for k in kham:
    ws2.append([k[0], datetime.combine(k[1], datetime.min.time()), k[2], k[3], k[4]])
luu(wb, "benh_nhan.xlsx")
B1, B2 = "benh_nhan.xlsx::BenhNhan", "benh_nhan.xlsx::KhamBenh"
TANG1[B1] = {"huong": ["NGANG"], "so_dong": 5, "kieu": {"Mã BN": "chu", "Ngày sinh": "ngay"},
             "o": [[4, "Ngày sinh", "1968-03-15"]], "nhan": ["nhieu_sheet", "ngay_excel"]}
TANG1[B2] = {"huong": ["NGANG"], "so_dong": 7,
             "kieu": {"Ngày khám": "ngay", "Chi phí": "so", "Tỷ lệ BHYT (%)": "so"},
             "o": [[0, "Tỷ lệ BHYT (%)", 80], [6, "Tỷ lệ BHYT (%)", 0]],
             "nhan": ["nhieu_sheet", "ngay_excel", "phan_tram"]}
Y = ["y_te"]
ma_bao = next(b[0] for b in benh_nhan if b[1] == "Trần Văn Bảo")
hoi("Tổng chi phí khám của bệnh nhân Trần Văn Bảo là bao nhiêu?",
    sum(k[3] for k in kham if k[0] == ma_bao), Y + ["ghep_2_bang", "tong_dieu_kien"], [B1, B2])
hoi("Bệnh nhân nào sinh năm 1992?", "Lê Thị Cúc", Y + ["ngay_excel", "tra_o"], [B1])
hoi("Có bao nhiêu lượt khám có tỷ lệ BHYT 80%?", sum(k[4] == "80%" for k in kham), Y + ["phan_tram", "dem_dieu_kien"], [B2])
hoi("Khoa Nội có bao nhiêu bệnh nhân?", sum(b[4] == "Nội" for b in benh_nhan), Y + ["dem_dieu_kien"], [B1])
hoi("Bệnh nhân lớn tuổi nhất tên gì?", min(benh_nhan, key=lambda b: b[2])[1], Y + ["ngay_excel", "max_tra_ten"], [B1])

# =====================================================================
# 7. TÀI CHÍNH – bảng DÀI (3000 dòng), ô trống, chữ dài
# =====================================================================
rng = random.Random(2026)
DM_THU = ["Lương", "Thưởng", "Bán hàng"]
DM_CHI = ["Ăn uống", "Đi lại", "Mua sắm", "Hóa đơn"]
GHI_CHU = ["Thanh toán qua chuyển khoản ngân hàng, đã đối soát với sao kê cuối tháng",
           "Chi phí phát sinh ngoài kế hoạch, cần trình duyệt bổ sung vào ngân sách quý",
           "Khoản thu định kỳ hằng tháng theo hợp đồng đã ký kết với đối tác"]
giao_dich = []
for _ in range(3000):
    ngay = date(2026, 1, 1) + timedelta(days=rng.randrange(181))
    if rng.random() < 0.25:
        loai, dm, tien = "Thu", rng.choice(DM_THU), rng.randrange(1000, 20000) * 1000
    else:
        loai, dm, tien = "Chi", rng.choice(DM_CHI), rng.randrange(20, 5000) * 1000
    giao_dich.append((ngay, loai, dm, tien, rng.choice(GHI_CHU) if rng.random() < 0.4 else None))
giao_dich.sort(key=lambda g: g[0])
wb = Workbook(); ws = wb.active; ws.title = "GiaoDich"
ws.append(["Ngày", "Loại", "Danh mục", "Số tiền", "Ghi chú"])
for g in giao_dich:
    ws.append([datetime.combine(g[0], datetime.min.time()), g[1], g[2], g[3], g[4]])
luu(wb, "thu_chi.xlsx")
B = "thu_chi.xlsx::GiaoDich"
TANG1[B] = {"huong": ["NGANG"], "so_dong": 3000,
            "kieu": {"Ngày": "ngay", "Loại": "chu", "Số tiền": "so", "Ghi chú": "chu"},
            "o": [[0, "Số tiền", giao_dich[0][3]]], "nhan": ["bang_dai", "o_trong", "chu_dai"]}
T = ["tai_chinh", "bang_dai"]
hoi("Tổng chi trong tháng 3/2026 là bao nhiêu?",
    sum(g[3] for g in giao_dich if g[1] == "Chi" and g[0].month == 3), T + ["ngay_excel", "tong_dieu_kien"], [B])
tong_dm = {dm: sum(g[3] for g in giao_dich if g[2] == dm) for dm in DM_CHI}
hoi("Danh mục chi nào tốn nhiều tiền nhất?", max(tong_dm, key=tong_dm.get), T + ["max_tra_ten"], [B])
hoi("Có bao nhiêu giao dịch thu?", sum(g[1] == "Thu" for g in giao_dich), T + ["dem_dieu_kien"], [B])

# =====================================================================
# 8. KHO – bảng RỘNG (33 cột), mã hàng có số 0 đầu
# =====================================================================
rng = random.Random(7)
hang = [("00123", "Bút bi"), ("00456", "Giấy A4"), ("01789", "Ghim bấm"),
        ("10234", "Băng keo"), ("00901", "Kéo"), ("20345", "Thước kẻ")]
ton = {}
for ma, _ in hang:
    x, ds = rng.randrange(100, 500), []
    for _ in range(31):
        x = max(0, x + rng.randrange(-30, 31))
        ds.append(x)
    ton[ma] = ds
wb = Workbook(); ws = wb.active; ws.title = "TonKho"
ws.append(["Mã hàng", "Tên hàng", *[f"Ngày {d}" for d in range(1, 32)]])
for ma, ten in hang:
    ws.append([ma, ten, *ton[ma]])
luu(wb, "ton_kho.xlsx")
B = "ton_kho.xlsx::TonKho"
TANG1[B] = {"huong": ["NGANG", "CHEO"], "so_dong": 6,
            "kieu": {"Mã hàng": "chu", "Tên hàng": "chu", "Ngày 1": "so", "Ngày 31": "so"},
            "o": [[0, "Mã hàng", "00123"], [3, "Mã hàng", "10234"]], "nhan": ["bang_rong", "ma_so_0"]}
K = ["kho", "bang_rong"]
hoi("Tồn kho ngày 15 của mã hàng 00123 là bao nhiêu?", ton["00123"][14], K + ["ma_so_0", "tra_o"], [B])
hoi("Mặt hàng nào có tồn kho ngày 31 thấp nhất?",
    next(t for m, t in hang if ton[m][30] == min(v[30] for v in ton.values())), K + ["max_tra_ten"], [B])

# =====================================================================
# 9. DOANH NGHIỆP – bảng DỌC chỉ 1 đối tượng, mã số thuế có số 0 đầu
# =====================================================================
wb = Workbook(); ws = wb.active; ws.title = "ThongTin"
for k, v in [("Tên công ty", "Công ty TNHH ABC Logistics"), ("Mã số thuế", "0312345678"),
             ("Năm thành lập", 2015), ("Số nhân viên", 120),
             ("Địa chỉ", "12 Nguyễn Huệ, Quận 1, TP.HCM")]:
    ws.append([k, v])
luu(wb, "cong_ty.xlsx")
B = "cong_ty.xlsx::ThongTin"
TANG1[B] = {"huong": ["DOC"], "so_dong": 1,
            "kieu": {"Mã số thuế": "chu", "Năm thành lập": "so", "Số nhân viên": "so"},
            "o": [[0, "Mã số thuế", "0312345678"]], "nhan": ["bang_doc", "bang_1_dong", "ma_so_0"]}
D = ["doanh_nghiep", "bang_doc", "bang_1_dong"]
hoi("Mã số thuế của công ty là gì?", "0312345678", D + ["ma_so_0", "tra_o"], [B])
hoi("Công ty có bao nhiêu nhân viên?", 120, D + ["tra_o"], [B])


# =====================================================================
# 11. BẪY cho quy tắc "tách bảng" và "xóa dòng cộng" (chỉ kiểm tra tầng 1) – những thứ KHÔNG được xóa/tách
# =====================================================================
wb = Workbook(); ws = wb.active; ws.title = "ChiTieu"
ws.append(["Chỉ tiêu", "Giá trị (triệu đồng)"])
for ct in [("Tổng doanh thu", 500), ("Chi phí", 300), ("Lợi nhuận", 200)]:
    ws.append(list(ct))
ws2 = wb.create_sheet("NhanSu")
ws2.append(["STT", "Họ tên", "Chức vụ", "Lương"])
ws2.append([1, "Phan Văn Tài", "Trưởng phòng", 25000000])
ws2.append([2, "Hà Thị Thu", "Nhân viên", 12000000])
ws2.append([3, "Kim Văn Lộc", "Cộng tác viên", 6000000])
ws2.append([])                                              # dòng trống lỡ tay GIỮA bảng
ws2.append([4, "Lò Thị Dung", "Nhân viên", 11000000])
ws2.append([5, "Mạc Văn Khoa", "Cộng tác viên", 5500000])
ws2.append([]); ws2.append([])
ws2.append([None, None, None, "Người lập biểu"])             # chữ ký cuối sheet
ws2.append([None, None, None, "Nguyễn Thị Kế Toán"])
luu(wb, "bao_cao_tai_chinh.xlsx")
TANG1["bao_cao_tai_chinh.xlsx::ChiTieu"] = {
    "huong": ["NGANG"], "so_dong": 3, "o": [[0, "Chỉ tiêu", "Tổng doanh thu"]],
    "nhan": ["bay_dong_tong_la_du_lieu"]}
TANG1["bao_cao_tai_chinh.xlsx::NhanSu"] = {
    "huong": ["NGANG"], "so_dong": 5, "kieu": {"Lương": "so"},
    "o": [[2, "Chức vụ", "Cộng tác viên"], [3, "Họ tên", "Lò Thị Dung"], [4, "Chức vụ", "Cộng tác viên"]],
    "nhan": ["bay_dong_trong_giua_bang", "bay_cong_tac_vien", "bay_chu_ky_cuoi_sheet"]}


# =====================================================================
# 10. Hai ví dụ "Văn – Lớp – Tuổi" (chỉ kiểm tra tầng 1): cùng nội dung, xoay khác chiều
# =====================================================================
wb = Workbook(); ws = wb.active; ws.title = "Ngang"
ws.append([None, "Lớp", "Tuổi"]); ws.append(["Văn", 12, 18])
ws2 = wb.create_sheet("Doc")
ws2.append([None, "Văn"]); ws2.append(["Lớp", 12]); ws2.append(["Tuổi", 18])
luu(wb, "vi_du_xoay.xlsx")
for sheet, huong in [("Ngang", "NGANG"), ("Doc", "DOC")]:
    TANG1[f"vi_du_xoay.xlsx::{sheet}"] = {
        "huong": [huong], "so_dong": 1, "cot": ["Tên", "Lớp", "Tuổi"],
        "o": [[0, "Tên", "Văn"], [0, "Tuổi", 18]], "nhan": ["bang_1_dong", "goc_trong"]}


# =====================================================================
# Ghi file
# =====================================================================
def dep(x):
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x)


queries, dap_an = [], {}
for i, (cau, da, nhan, bang) in enumerate(CAU_HOI, 1):
    queries.append({"id": i, "question": cau})
    dap_an[str(i)] = {"answer": dep(da), "nhan": nhan, "bang": bang}

for ten, du_lieu in [("queries.json", queries), ("dap_an.json", dap_an), ("tang1.json", TANG1)]:
    with open(os.path.join(GOC, ten), "w", encoding="utf-8") as f:
        json.dump(du_lieu, f, ensure_ascii=False, indent=2)

print(f"Đã tạo {len(os.listdir(XLSX))} file Excel ({len(TANG1)} bảng), {len(queries)} câu hỏi")

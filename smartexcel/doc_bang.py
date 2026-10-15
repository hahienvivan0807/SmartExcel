"""BƯỚC 1 – Đọc & làm sạch bảng.

Mỗi sheet đi qua các bước:
    1. Đọc ô, điền lại giá trị cho ô bị gộp (merge)
    2. Bỏ dòng/cột trống, cắt dòng tiêu đề báo cáo ở trên ("TRƯỜNG THPT ABC")
    3. Xác định kiểu bảng NGANG / DOC / CHEO (hỏi AI góc 6×6, thuật toán dự phòng)
    4. Bảng DOC → xoay lại thành ngang
    5. Đặt tên cột (gộp tiêu đề 2 tầng), xóa dòng "Tổng cộng", chuyển số
Kết quả: mỗi dòng = 1 đối tượng, mỗi cột = 1 thuộc tính.
"""
import json
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .log import log
from .van_ban import bo_dau

DUOI_FILE = (".xlsx", ".xlsm")

# Từ hay gặp trong tên thuộc tính (không dấu) – dùng cho thuật toán đoán hướng
TU_THUOC_TINH = {
    "ten", "ho", "lop", "tuoi", "ngay", "nam", "thang", "quy", "diem", "ma",
    "gia", "so", "luong", "tong", "gioi", "tinh", "dia", "chi", "sdt", "email",
    "chuc", "vu", "phong", "ban", "don", "vi", "loai", "nhom", "san", "pham",
    "khach", "hang", "doanh", "thu", "ghi", "chu", "trang", "thai", "stt",
    # logistics / kho
    "kho", "xuat", "nhap", "ton", "tuyen", "trong", "cuoc", "phi", "gui", "nhan",
    "buu", "dien", "thoai", "tai", "xe",
    # y tế
    "benh", "khoa", "kham", "chan", "doan", "thuoc", "bhyt",
    # tài chính / doanh nghiệp
    "tien", "muc", "danh", "thue", "von", "lai", "huyen", "xa", "quoc", "gia",
    # KHÔNG thêm từ dễ trùng tên người/tên công ty: van (Văn), le (Lê), cong/ty (Công ty TNHH…)
}

TU_DONG_TONG = {"tong", "tong cong", "cong", "total", "grand total", "tong so"}


@dataclass
class Bang:
    ten: str          # "file.xlsx::Sheet1"
    df: pd.DataFrame  # bảng đã làm sạch
    huong: str        # NGANG / DOC / CHEO
    mo_ta: str        # câu mô tả ngắn (AI đưa, có thể rỗng)
    tieu_de: str = ""  # dòng tiêu đề ngay trên bảng, vd "BẢNG NHẬP KHO"
    tom_tat: str = ""                          # Bước 2 điền
    tu_khoa: set = field(default_factory=set)  # Bước 2 điền – dùng để lọc bảng ở Bước 3


# ============================================================
# Đọc file
# ============================================================

def doc_tat_ca(thu_muc, ai=None):
    """Đọc mọi sheet của mọi file Excel trong thư mục (kể cả thư mục con)."""
    cac_bang = []
    for goc, _, files in os.walk(thu_muc):
        for ten_file in sorted(files):
            if ten_file.startswith("~$") or not ten_file.lower().endswith(DUOI_FILE):
                continue
            duong_dan = os.path.join(goc, ten_file)
            ten_tuong_doi = os.path.relpath(duong_dan, thu_muc).replace("\\", "/")
            try:
                wb = load_workbook(duong_dan, data_only=True)
            except Exception as e:
                log(f"Không mở được {ten_tuong_doi}: {e}")
                continue
            for ws in wb.worksheets:
                ten_sheet = f"{ten_tuong_doi}::{ws.title}"
                try:
                    cac_khoi = tach_bang(_doc_sheet(ws))
                except Exception as e:
                    log(f"Lỗi đọc {ten_sheet}: {e}")
                    continue
                if len(cac_khoi) > 1:
                    log(f"[{ten_sheet}] có {len(cac_khoi)} bảng trong 1 sheet")
                for i, khoi in enumerate(cac_khoi, 1):
                    # 1 bảng → "file::Sheet"; nhiều bảng → "file::Sheet#1", "file::Sheet#2"…
                    ten = ten_sheet if len(cac_khoi) == 1 else f"{ten_sheet}#{i}"
                    try:
                        bang = lam_sach_sheet(ten, khoi, ai)
                        if bang is not None:
                            cac_bang.append(bang)
                    except Exception as e:
                        log(f"Lỗi xử lý {ten}: {e}")
    return cac_bang


def _doc_sheet(ws):
    """Đọc sheet thành DataFrame thô (chưa có tiêu đề), điền lại ô gộp."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for vung in ws.merged_cells.ranges:
        v = grid[vung.min_row - 1][vung.min_col - 1] if grid else None
        for r in range(vung.min_row - 1, min(vung.max_row, len(grid))):
            for c in range(vung.min_col - 1, min(vung.max_col, len(grid[r]))):
                grid[r][c] = v
    raw = pd.DataFrame(grid, dtype=object)
    return raw.map(_lam_sach_o)


def _lam_sach_o(x):
    if isinstance(x, str):
        x = re.sub(r"\s+", " ", x).strip()
        return x if x else None
    return x


# ============================================================
# Làm sạch 1 sheet
# ============================================================

def lam_sach_sheet(ten, raw, ai=None):
    raw, cac_tieu_de = _cat_tieu_de_bao_cao(_bo_trong(raw))
    if raw.empty:
        return None

    huong, mo_ta = xac_dinh_huong(ten, raw, ai)
    if huong == "DOC":
        raw = _bo_trong(raw.T)

    df = _dung_bang(raw)
    df = _xoa_dong_tong(df)
    df = _chuyen_kieu(df)
    if df.empty:
        return None
    tieu_de = cac_tieu_de[-1] if cac_tieu_de else ""   # dòng gần bảng nhất, vd "BẢNG NHẬP KHO"
    return Bang(ten=ten, df=df, huong=huong, mo_ta=mo_ta, tieu_de=tieu_de)


def _bo_trong(raw):
    raw = raw.dropna(how="all").dropna(axis=1, how="all")
    raw.index = range(len(raw))
    raw.columns = range(raw.shape[1])
    return raw


def _cat_tieu_de_bao_cao(raw):
    """Bỏ các dòng đầu kiểu 'TRƯỜNG THPT ABC': chỉ 1 giá trị, hoặc trống quá nửa.
    Trả về (bảng còn lại, [chữ của các dòng đã cắt])."""
    so_cot = raw.shape[1]
    cac_tieu_de = []
    if so_cot < 3:
        return raw, cac_tieu_de
    while len(raw) > 1:
        o = raw.iloc[0].dropna()
        if o.nunique() <= 1 or len(o) < so_cot * 0.5:
            cac_tieu_de.append(" ".join(dict.fromkeys(str(x) for x in o)))
            raw = raw.iloc[1:]
        else:
            break
    raw.index = range(len(raw))
    return raw, cac_tieu_de


# ============================================================
# Tách nhiều bảng trong 1 sheet
# ============================================================

def tach_bang(raw):
    """Tách sheet thành các bảng, ngăn cách bởi dòng trống.

    Không phải dòng trống nào cũng ngăn 2 bảng:
      - Khối chỉ có dòng tiêu đề / ghi chú (không đủ 2 dòng × 2 cột) → không phải bảng:
        gắn vào khối sau làm tiêu đề (vd "TRƯỜNG THPT ABC" ↵ trống ↵ bảng điểm);
        nằm cuối sheet (vd "Người lập biểu") → bỏ.
      - Cách đúng 1 dòng trống, không có tiêu đề riêng, dòng đầu không giống tên cột
        → chỉ là dòng trống lỡ tay GIỮA bảng → nối lại vào bảng trước.
    """
    trong = raw.isna().all(axis=1).tolist()
    khoi, bat_dau, dem_trong, trong_truoc = [], None, 0, 0   # khối = (dòng đầu, dòng cuối, số dòng trống phía trước)
    for i, t in enumerate(trong):
        if t:
            if bat_dau is not None:
                khoi.append((bat_dau, i - 1, trong_truoc))
                bat_dau = None
            dem_trong += 1
        elif bat_dau is None:
            bat_dau, trong_truoc, dem_trong = i, dem_trong, 0
    if bat_dau is not None:
        khoi.append((bat_dau, len(trong) - 1, trong_truoc))
    if len(khoi) <= 1:
        return [raw]

    ket_qua, tien_to = [], None
    for dau, cuoi, so_trong in khoi:
        phan = raw.iloc[dau:cuoi + 1]
        if tien_to is not None:
            phan, tien_to = pd.concat([tien_to, phan]), None
        con_lai, cac_tieu_de = _cat_tieu_de_bao_cao(_bo_trong(phan))
        if len(con_lai) < 2 or con_lai.shape[1] < 2:      # chỉ là tiêu đề / ghi chú
            tien_to = phan
            continue
        if ket_qua and so_trong == 1 and not cac_tieu_de and not _dong_dau_la_ten_cot(con_lai):
            ket_qua[-1] = pd.concat([ket_qua[-1], phan])  # dòng trống lẻ giữa bảng → nối lại
            continue
        ket_qua.append(phan)
    return ket_qua or [raw]


def _dong_dau_la_ten_cot(raw):
    """Dòng đầu là chữ mà dòng 2 cùng cột là số/ngày → dòng đầu là tên cột (bảng mới)."""
    if len(raw) < 2:
        return False
    d1, d2 = raw.iloc[0].tolist(), raw.iloc[1].tolist()
    return any(_kieu(a) == "chu" and _kieu(b) in ("so", "ngay") for a, b in zip(d1, d2))


def _dung_bang(raw):
    """Dòng đầu (hoặc 2 dòng đầu nếu tiêu đề 2 tầng) thành tên cột."""
    so_dong_tieu_de = 2 if _tieu_de_2_tang(raw) else 1
    tieu_de = raw.iloc[:so_dong_tieu_de]

    ten_cot, da_co = [], set()
    for c in raw.columns:
        phan = []
        for v in tieu_de[c]:
            if v is not None and not pd.isna(v) and str(v) not in phan:
                phan.append(str(v).strip())
        ten = " - ".join(phan) or ("Tên" if c == 0 else f"Cột {c + 1}")
        goc, i = ten, 2
        while ten in da_co:          # tên trùng → thêm số
            ten = f"{goc} ({i})"
            i += 1
        da_co.add(ten)
        ten_cot.append(ten)

    df = raw.iloc[so_dong_tieu_de:].copy()
    df.columns = ten_cot
    df = df.dropna(how="all").infer_objects()
    df.index = range(len(df))
    return df


def _tieu_de_2_tang(raw):
    """Dòng 1 có ô gộp ngang (giá trị lặp kề nhau) và dòng 2 cũng là chữ → 2 tầng."""
    if len(raw) < 3:
        return False
    d1, d2 = list(raw.iloc[0]), list(raw.iloc[1])
    lap_ke_nhau = any(d1[i] is not None and d1[i] == d1[i + 1] for i in range(len(d1) - 1))
    o_d2 = [x for x in d2 if x is not None and not pd.isna(x)]
    la_chu = [x for x in o_d2 if _kieu(x) == "chu"]
    return (lap_ke_nhau
            and len(o_d2) >= len(d2) * 0.5
            and len(la_chu) >= len(o_d2) * 0.6
            and d1 != d2)


# Nhãn dòng cộng nhóm: "Cộng Miền Bắc", "Tổng cộng", "Subtotal"…
# KHÔNG gồm "Tổng …" chung chung vì dễ là dữ liệu thật ("Tổng doanh thu | 500" trong bảng chỉ tiêu)
_NHAN_CONG = ("cong", "tong cong", "subtotal", "sub total", "total", "grand total")


def _la_nhan_cong(v):
    s = bo_dau(v)
    return s == "tong" or any(s == t or s.startswith(t + " ") for t in _NHAN_CONG)


def _xoa_dong_tong(df):
    """Xóa dòng tổng/cộng nhóm – ở CUỐI bảng lẫn GIỮA bảng – để khi tính tổng không bị cộng 2 lần.

    Dòng bị xóa khi có ô nhãn ("Cộng Miền Bắc", "Tổng cộng"…) và:
      - là 1 trong 2 dòng cuối, ô ghi đúng "Tổng/Tổng cộng/Cộng/Total…"; hoặc
      - nhãn là ô chữ DUY NHẤT của dòng (ô chữ khác nếu có chỉ lặp tên nhóm) và dòng có số.
    "Cộng tác viên" trong dòng có đủ tên, chức vụ… → còn ô chữ khác → giữ lại.
    """
    if len(df) < 2:
        return df
    cuoi = set(df.index[-2:])
    bo = []
    for i in df.index:
        dong = [v for v in df.loc[i] if v is not None and not (isinstance(v, float) and pd.isna(v))]
        chu = [v for v in dong if isinstance(v, str) and thanh_so(v) is None]
        nhan = [v for v in chu if _la_nhan_cong(v)]
        if not nhan:
            continue
        if i in cuoi and any(bo_dau(v) in TU_DONG_TONG or bo_dau(v).startswith("tong cong") for v in nhan):
            bo.append(i)
            continue
        khac = [v for v in chu if v not in nhan]
        co_so = any(thanh_so(v) is not None for v in dong)
        if co_so and (not khac or (len(khac) == 1 and bo_dau(khac[0]) in bo_dau(nhan[0]))):
            bo.append(i)
    return df.drop(index=bo).reset_index(drop=True)


# ============================================================
# Chuyển kiểu dữ liệu: mã (giữ chữ) / ngày / số (bỏ đơn vị)
# ============================================================

# Đơn vị hay gặp → tên chuẩn ghi vào tên cột
DON_VI = {
    "kg": "kg", "g": "g", "tấn": "tấn", "km": "km", "m": "m", "cm": "cm", "mm": "mm",
    "l": "lít", "lít": "lít", "ml": "ml", "%": "%",
    "đ": "đ", "vnđ": "đ", "vnd": "đ", "đồng": "đ", "usd": "USD",
    "cái": "cái", "chiếc": "chiếc", "thùng": "thùng", "hộp": "hộp", "kiện": "kiện",
    "ngày": "ngày", "giờ": "giờ", "phút": "phút", "tháng": "tháng", "năm": "năm", "người": "người",
}
_MAU_SO_DON_VI = re.compile(
    r"(-?[\d.,]+)\s*(" + "|".join(sorted(map(re.escape, DON_VI), key=len, reverse=True)) + r")?",
    re.IGNORECASE)
_MAU_NGAY_VN = re.compile(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})")    # 15/10/2026
_MAU_NGAY_ISO = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})")             # 2026-10-15


def la_ma_so_0(x):
    """'0901234567', '070000', '00123' → mã (số 0 đầu), KHÔNG được đổi thành số."""
    return isinstance(x, str) and re.fullmatch(r"0\d+", x.strip()) is not None


def _doc_phan_so(s):
    """Phần số: '1.200.000' → 1200000, '12,5' → 12.5"""
    if re.fullmatch(r"-?\d{1,3}([.,]\d{3})+", s):        # dấu phân cách hàng nghìn
        s = re.sub(r"[.,]", "", s)
    elif re.fullmatch(r"-?\d+,\d+", s):                  # dấu phẩy thập phân
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def tach_so_don_vi(x):
    """'12,5 kg' → (12.5, 'kg'), '35.000 đ' → (35000, 'đ'), '80%' → (80, '%'), '1.200' → (1200, None).
    Không phải số (hoặc là mã có số 0 đầu) → (None, None)."""
    if isinstance(x, (bool, np.bool_)):
        return None, None
    if isinstance(x, (int, float, np.integer, np.floating)):   # gồm cả số numpy sau khi pandas đọc
        return (None, None) if pd.isna(x) else (x, None)
    if not isinstance(x, str) or la_ma_so_0(x):
        return None, None
    m = _MAU_SO_DON_VI.fullmatch(x.strip())
    if not m:
        return None, None
    so = _doc_phan_so(m.group(1))
    if so is None:
        return None, None
    return so, DON_VI[m.group(2).lower()] if m.group(2) else None


def thanh_so(x):
    return tach_so_don_vi(x)[0]


def doc_ngay(x):
    """'15/10/2026', '15-10-2026', '2026-10-15' → datetime. Không phải ngày → None."""
    if isinstance(x, datetime):
        return x
    if isinstance(x, date):
        return datetime(x.year, x.month, x.day)
    if not isinstance(x, str):
        return None
    s = x.strip()
    try:
        if m := _MAU_NGAY_VN.fullmatch(s):
            return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        if m := _MAU_NGAY_ISO.fullmatch(s):
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:      # vd 31/02/2026
        return None
    return None


def _chuyen_kieu(df):
    """Với mỗi cột, theo thứ tự:
    1. Có ô là mã số 0 đầu ('0901234567') → cả cột giữ dạng chữ (là cột mã)
    2. ≥ 80% ô là ngày (Excel hoặc '15/10/2026') → cột ngày
    3. ≥ 80% ô là số, cùng 1 đơn vị → cột số; đơn vị ghi vào tên cột: 'Trọng lượng (kg)'
    """
    doi_ten = {}
    for col in df.columns:
        o = df[col].dropna()
        if o.empty or pd.api.types.is_datetime64_any_dtype(df[col]):
            continue

        if any(la_ma_so_0(x) for x in o):
            df[col] = df[col].map(lambda x: x if x is None or isinstance(x, str) or pd.isna(x) else str(x))
            continue

        if o.map(doc_ngay).notna().mean() >= 0.8:
            df[col] = pd.to_datetime(df[col].map(doc_ngay), errors="coerce")
            continue

        tach = o.map(tach_so_don_vi)
        so = tach.map(lambda t: t[0])
        cac_don_vi = {t[1] for t in tach if t[1]}
        if so.notna().mean() < 0.8 or len(cac_don_vi) > 1:   # ít số, hoặc lẫn đơn vị (kg + tấn) → giữ chữ
            continue
        cot = pd.to_numeric(df[col].map(thanh_so), errors="coerce")
        if (cot.dropna() % 1 == 0).all():   # toàn số nguyên → bỏ ".0"
            cot = cot.astype("Int64")
        df[col] = cot
        if cac_don_vi:
            dv = cac_don_vi.pop()
            if bo_dau(dv) not in bo_dau(col):
                doi_ten[col] = f"{col} ({dv})"
    return df.rename(columns=doi_ten)


# ============================================================
# Xác định hướng bảng: AI trước, thuật toán dự phòng
# ============================================================

def xac_dinh_huong(ten, raw, ai=None):
    """Trả về (huong, mo_ta) với huong ∈ NGANG / DOC / CHEO."""
    if raw.shape[0] < 2 or raw.shape[1] < 2:
        return "NGANG", ""

    huong_tt = doan_huong_thuat_toan(raw)
    kq_ai = phan_loai_bang_ai(raw, ai) if ai is not None else None

    if kq_ai is None:
        huong = huong_tt if huong_tt != "KHONG_CHAC" else "NGANG"
        log(f"[{ten}] hướng = {huong} (thuật toán)")
        return huong, ""

    huong_ai = kq_ai["huong"]
    so_sanh = "NGANG" if huong_ai == "CHEO" else huong_ai
    if huong_tt != "KHONG_CHAC" and huong_tt != so_sanh:
        log(f"⚠️ [{ten}] AI và thuật toán bất đồng: AI={huong_ai}, thuật toán={huong_tt}")
    log(f"[{ten}] hướng = {huong_ai} (AI)")
    return huong_ai, kq_ai.get("mo_ta", "")


def _kieu(x):
    """Loại của 1 ô: 'so', 'ngay', 'chu' hoặc None (trống)."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    if isinstance(x, (datetime, date)):
        return "ngay"
    if thanh_so(x) is not None:
        return "so"
    return "chu"


def _do_dong_nhat(cac_day):
    diem = []
    for day in cac_day:
        k = [t for t in map(_kieu, day) if t]
        if len(k) >= 2:
            diem.append(max(k.count(t) for t in set(k)) / len(k))
    return sum(diem) / len(diem) if diem else 0


def _giong_thuoc_tinh(gia_tri):
    """% ô chữ có TỪ ĐẦU TIÊN là từ thuộc tính.
    Chỉ xét từ đầu: 'Mã số thuế', 'Ngày gửi' mở đầu bằng từ loại;
    tên người 'Nguyễn Văn An' mở đầu bằng họ nên 'Văn' ở giữa không bị tính."""
    chu = [v for v in gia_tri if _kieu(v) == "chu"]
    if not chu:
        return 0
    tu_dau = [re.findall(r"[^\W_]+", bo_dau(v))[:1] for v in chu]
    return sum(bool(t) and t[0] in TU_THUOC_TINH for t in tu_dau) / len(chu)


def doan_huong_thuat_toan(raw):
    """Bỏ phiếu bằng 3 dấu hiệu → 'NGANG', 'DOC' hoặc 'KHONG_CHAC'."""
    g = raw.values.tolist()
    dong_dau = g[0][1:]
    cot_dau = [r[0] for r in g[1:]]
    than = [r[1:] for r in g[1:]]
    ngang = doc = 0

    # 1. Từ ngữ giống tên thuộc tính (2 phiếu)
    a, b = _giong_thuoc_tinh(dong_dau), _giong_thuoc_tinh(cot_dau)
    if a > b:
        ngang += 2
    elif b > a:
        doc += 2

    # 2. Dữ liệu cùng loại theo cột hay theo dòng (1 phiếu)
    if len(than) >= 2 and len(than[0]) >= 2:
        theo_cot, theo_dong = _do_dong_nhat(zip(*than)), _do_dong_nhat(than)
        if theo_cot > theo_dong + 0.1:
            ngang += 1
        elif theo_dong > theo_cot + 0.1:
            doc += 1

    so_dong, so_cot = raw.shape
    if so_cot == 2:
        # 3a. Bảng 2 cột: cột 2 lẫn nhiều loại (chữ, số, ngày…) → danh sách "thuộc tính | giá trị" → dọc
        #     (bỏ dấu hiệu hình dạng vì danh sách thuộc tính luôn "dài")
        loai = {t for t in map(_kieu, [r[1] for r in g[1:]]) if t}
        if len(loai) >= 2 and so_dong >= 3:
            doc += 1
    # 3b. Hình dạng: dài → ngang, rộng → dọc (1 phiếu)
    elif so_dong > so_cot * 1.5:
        ngang += 1
    elif so_cot > so_dong * 1.5:
        doc += 1

    if ngang > doc:
        return "NGANG"
    if doc > ngang:
        return "DOC"
    return "KHONG_CHAC"


PROMPT_PHAN_LOAI = """Phân loại bảng Excel. Trả về JSON đúng mẫu, không giải thích:
{"huong": "NGANG" | "DOC" | "CHEO", "mo_ta": "<1 câu: mỗi dòng/cột/ô là gì>"}

NGANG = dòng đầu là tên thuộc tính, mỗi dòng bên dưới là 1 đối tượng
DOC   = cột đầu là tên thuộc tính, mỗi cột bên phải là 1 đối tượng
CHEO  = dòng đầu và cột đầu đều là tiêu đề, ô bên trong cùng 1 loại số liệu

Ví dụ 1:
 | Lớp | Tuổi
Văn | 12 | 18
An | 11 | 17
→ {"huong":"NGANG","mo_ta":"mỗi dòng là 1 học sinh"}

Ví dụ 2:
 | Văn | An
Lớp | 12 | 11
Tuổi | 18 | 17
→ {"huong":"DOC","mo_ta":"mỗi cột là 1 học sinh"}

Ví dụ 3:
 | Toán | Văn
An | 8 | 7
Bình | 6 | 9
→ {"huong":"CHEO","mo_ta":"dòng là học sinh, cột là môn, ô là điểm"}"""


def phan_loai_bang_ai(raw, ai):
    """Gửi góc trên-trái 6×6 cho AI. Trả về dict hoặc None nếu thất bại."""
    goc = raw.iloc[:6, :6].map(lambda x: "" if x is None or pd.isna(x) else str(x)[:20])
    xem_truoc = "\n".join(" | ".join(r) for r in goc.values.tolist())
    noi_dung = f"Bảng có {raw.shape[0]} dòng × {raw.shape[1]} cột. Góc trên-trái:\n{xem_truoc}"
    try:
        tra_loi = ai.hoi([
            {"role": "system", "content": PROMPT_PHAN_LOAI},
            {"role": "user", "content": noi_dung},
        ], max_tokens=150)
        kq = json.loads(tra_loi[tra_loi.find("{"): tra_loi.rfind("}") + 1])
        kq["huong"] = str(kq.get("huong", "")).upper()
        if kq["huong"] in ("NGANG", "DOC", "CHEO"):
            return kq
    except Exception as e:
        log(f"AI phân loại bảng lỗi: {e}")
    return None
